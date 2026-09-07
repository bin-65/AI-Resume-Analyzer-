"""Resume file extraction. This module knows nothing about AI or UI concerns."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path

from docx import Document
from pypdf import PdfReader


class ResumeParseError(Exception):
    """Raised when a supported resume cannot be read."""


def supported_file_types() -> list[str]:
    return ["pdf", "docx"]


def _extract_pdf(file_bytes: bytes) -> str:
    reader = PdfReader(BytesIO(file_bytes))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def _extract_docx(file_bytes: bytes) -> str:
    document = Document(BytesIO(file_bytes))
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    table_cells = [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    return "\n".join(paragraphs + table_cells)


def extract_resume_text(file_bytes: bytes, filename: str) -> str:
    """Extract normalized text from a PDF or DOCX resume."""
    suffix = Path(filename).suffix.lower()
    try:
        if suffix == ".pdf":
            text = _extract_pdf(file_bytes)
        elif suffix == ".docx":
            text = _extract_docx(file_bytes)
        else:
            raise ResumeParseError("Please upload a PDF or DOCX file.")
    except ResumeParseError:
        raise
    except Exception as error:
        raise ResumeParseError("The file is damaged, password-protected, or unsupported.") from error

    cleaned = "\n".join(line.strip() for line in text.splitlines() if line.strip())
    if len(cleaned) < 30:
        raise ResumeParseError("Very little text was found. Use a text-based PDF or DOCX resume.")
    return cleaned
