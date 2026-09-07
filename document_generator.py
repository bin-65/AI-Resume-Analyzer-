"""Create downloadable DOCX and PDF resumes from structured, truthful resume data."""

from __future__ import annotations

from io import BytesIO
from typing import Any, Iterable

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
from reportlab.pdfbase.pdfmetrics import stringWidth


BLUE = "1F4E79"


def _text(value: Any) -> str:
    return str(value).strip() if value else ""


def _items(value: Any) -> list[str]:
    if not isinstance(value, list):
        return []
    return [_text(item) for item in value if _text(item)]


def _records(value: Any) -> list[dict[str, Any]]:
    return [item for item in value if isinstance(item, dict)] if isinstance(value, list) else []


def _contact_line(details: dict[str, Any]) -> str:
    return " | ".join(
        value for value in (_text(details.get(field)) for field in ("email", "phone", "location", "linkedin")) if value
    )


def _docx_heading(document: Document, heading: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(10)
    paragraph.paragraph_format.space_after = Pt(3)
    run = paragraph.add_run(heading.upper())
    run.bold = True
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(31, 78, 121)


def _docx_bullets(document: Document, bullets: Iterable[str]) -> None:
    for bullet in bullets:
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(1)
        paragraph.add_run(bullet)


def create_docx(resume: dict[str, Any]) -> bytes:
    """Return a clean one-column Word resume as bytes."""
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)

    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(3)

    details = resume.get("personal_details", {}) if isinstance(resume.get("personal_details"), dict) else {}
    name = _text(details.get("name")) or "Resume"
    title = _text(resume.get("headline"))
    name_paragraph = document.add_paragraph()
    name_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name_run = name_paragraph.add_run(name)
    name_run.bold = True
    name_run.font.size = Pt(20)
    name_run.font.color.rgb = RGBColor(31, 78, 121)
    if title:
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run(title).italic = True
    contact = _contact_line(details)
    if contact:
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run(contact).font.size = Pt(9)

    summary = _text(resume.get("professional_summary"))
    if summary:
        _docx_heading(document, "Professional Summary")
        document.add_paragraph(summary)

    skills = _items(resume.get("skills"))
    if skills:
        _docx_heading(document, "Skills")
        document.add_paragraph(" • ".join(skills))

    experience = _records(resume.get("experience"))
    if experience:
        _docx_heading(document, "Experience")
        for item in experience:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_before = Pt(3)
            role = " | ".join(value for value in (_text(item.get("title")), _text(item.get("company"))) if value)
            dates = _text(item.get("dates"))
            run = paragraph.add_run(role)
            run.bold = True
            if dates:
                paragraph.add_run(f"  ({dates})")
            _docx_bullets(document, _items(item.get("bullets")))

    education = _records(resume.get("education"))
    if education:
        _docx_heading(document, "Education")
        for item in education:
            paragraph = document.add_paragraph()
            degree = " | ".join(value for value in (_text(item.get("degree")), _text(item.get("institution"))) if value)
            paragraph.add_run(degree).bold = True
            dates = _text(item.get("dates"))
            if dates:
                paragraph.add_run(f"  ({dates})")
            details_text = _text(item.get("details"))
            if details_text:
                document.add_paragraph(details_text)

    projects = _records(resume.get("projects"))
    if projects:
        _docx_heading(document, "Projects")
        for item in projects:
            paragraph = document.add_paragraph()
            paragraph.add_run(_text(item.get("name"))).bold = True
            details_text = _text(item.get("details"))
            if details_text:
                paragraph.add_run(f" - {details_text}")
            _docx_bullets(document, _items(item.get("bullets")))

    certifications = _items(resume.get("certifications"))
    if certifications:
        _docx_heading(document, "Certifications")
        _docx_bullets(document, certifications)

    for section_data in _records(resume.get("additional_sections")):
        title = _text(section_data.get("title"))
        items = _items(section_data.get("items"))
        if title and items:
            _docx_heading(document, title)
            _docx_bullets(document, items)

    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _pdf_styles() -> dict[str, ParagraphStyle]:
    styles = getSampleStyleSheet()
    return {
        "name": ParagraphStyle("ResumeName", parent=styles["Title"], fontName="Helvetica-Bold", fontSize=20, leading=24, textColor=colors.HexColor(f"#{BLUE}"), alignment=TA_CENTER, spaceAfter=3),
        "center": ParagraphStyle("ResumeCenter", parent=styles["Normal"], fontName="Helvetica", fontSize=9, leading=11, alignment=TA_CENTER, spaceAfter=2),
        "body": ParagraphStyle("ResumeBody", parent=styles["Normal"], fontName="Helvetica", fontSize=9.5, leading=13, alignment=TA_LEFT, spaceAfter=3),
        "heading": ParagraphStyle("ResumeHeading", parent=styles["Heading2"], fontName="Helvetica-Bold", fontSize=10, leading=13, textColor=colors.HexColor(f"#{BLUE}"), spaceBefore=9, spaceAfter=3),
        "role": ParagraphStyle("ResumeRole", parent=styles["Normal"], fontName="Helvetica-Bold", fontSize=9.5, leading=13, spaceBefore=2, spaceAfter=1),
        "bullet": ParagraphStyle("ResumeBullet", parent=styles["Normal"], fontName="Helvetica", fontSize=9.5, leading=12, leftIndent=13, firstLineIndent=-7, spaceAfter=1),
    }


def _paragraph(value: str, style: ParagraphStyle) -> Paragraph:
    safe = value.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    return Paragraph(safe, style)


def _pdf_heading(story: list[Any], styles: dict[str, ParagraphStyle], title: str) -> None:
    story.append(_paragraph(title.upper(), styles["heading"]))


def _pdf_bullets(story: list[Any], styles: dict[str, ParagraphStyle], bullets: Iterable[str]) -> None:
    for bullet in bullets:
        story.append(_paragraph(f"• {bullet}", styles["bullet"]))


def create_pdf(resume: dict[str, Any]) -> bytes:
    """Return a matching ATS-friendly PDF resume as bytes."""
    output = BytesIO()
    document = SimpleDocTemplate(output, pagesize=A4, leftMargin=0.65 * inch, rightMargin=0.65 * inch, topMargin=0.55 * inch, bottomMargin=0.55 * inch)
    styles = _pdf_styles()
    story: list[Any] = []
    details = resume.get("personal_details", {}) if isinstance(resume.get("personal_details"), dict) else {}
    story.append(_paragraph(_text(details.get("name")) or "Resume", styles["name"]))
    if _text(resume.get("headline")):
        story.append(_paragraph(_text(resume.get("headline")), styles["center"]))
    contact = _contact_line(details)
    if contact:
        story.append(_paragraph(contact, styles["center"]))

    summary = _text(resume.get("professional_summary"))
    if summary:
        _pdf_heading(story, styles, "Professional Summary")
        story.append(_paragraph(summary, styles["body"]))
    skills = _items(resume.get("skills"))
    if skills:
        _pdf_heading(story, styles, "Skills")
        story.append(_paragraph(" • ".join(skills), styles["body"]))
    experience = _records(resume.get("experience"))
    if experience:
        _pdf_heading(story, styles, "Experience")
        for item in experience:
            role = " | ".join(value for value in (_text(item.get("title")), _text(item.get("company"))) if value)
            dates = _text(item.get("dates"))
            story.append(_paragraph(f"{role}{f' ({dates})' if dates else ''}", styles["role"]))
            _pdf_bullets(story, styles, _items(item.get("bullets")))
    education = _records(resume.get("education"))
    if education:
        _pdf_heading(story, styles, "Education")
        for item in education:
            value = " | ".join(part for part in (_text(item.get("degree")), _text(item.get("institution"))) if part)
            dates = _text(item.get("dates"))
            story.append(_paragraph(f"{value}{f' ({dates})' if dates else ''}", styles["role"]))
            if _text(item.get("details")):
                story.append(_paragraph(_text(item.get("details")), styles["body"]))
    projects = _records(resume.get("projects"))
    if projects:
        _pdf_heading(story, styles, "Projects")
        for item in projects:
            value = _text(item.get("name"))
            if _text(item.get("details")):
                value = f"{value} - {_text(item.get('details'))}"
            story.append(_paragraph(value, styles["role"]))
            _pdf_bullets(story, styles, _items(item.get("bullets")))
    certifications = _items(resume.get("certifications"))
    if certifications:
        _pdf_heading(story, styles, "Certifications")
        _pdf_bullets(story, styles, certifications)
    for section_data in _records(resume.get("additional_sections")):
        title = _text(section_data.get("title"))
        items = _items(section_data.get("items"))
        if title and items:
            _pdf_heading(story, styles, title)
            _pdf_bullets(story, styles, items)

    document.build(story)
    return output.getvalue()
